#!/usr/bin/env python3
"""Check Grouped Div (J in file2, I in file1) for IPSS GAP separation"""
import openpyxl
from collections import defaultdict

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'

# File #2: check J(Grouped Div) -> K(Division) mapping
print('=== File #2: Grouped Div(J) -> Division(K) for GAP ===')
f2 = BASE + r'\2) 영업DATA_Raw_v1.3_20260608_openclaw 작업용.xlsx'
wb2 = openpyxl.load_workbook(f2, read_only=True, data_only=True)
ws2 = wb2['Raw_20260608']

gd2_div = defaultdict(lambda: defaultdict(int))
seasons = {'26 SPR','26 SUM','26 FAL','26 HOL'}

for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0: continue
    v = list(row)
    season = str(v[12] or '').strip()
    if season not in seasons: continue
    q = v[26] or 0
    if not q: continue
    gd = str(v[9] or '').strip()      # J열 = Grouped Div (index 9)
    div = str(v[10] or '').strip()    # K열 = Division (index 10)
    if 'GAP' in gd.upper() or 'GAP' in div.upper() or 'IPSS' in gd.upper() or 'IPSS' in div.upper():
        gd2_div[gd][div] += q

wb2.close()
for gd in sorted(gd2_div.keys()):
    for div, q in sorted(gd2_div[gd].items(), key=lambda x: -x[1]):
        print(f'  Grouped Div={gd:20s} -> Division={div:20s} qty={int(q):>10,}')

# File #1: check what Grouped Div exists
print('\n=== File #1: Grouped Div(I) values for GAP ===')
f1 = BASE + r'\1) 브랜드별 직접원가_raw_v1.0_20260629.xlsx'
wb1 = openpyxl.load_workbook(f1, read_only=True, data_only=True)
ws1 = wb1['Raw_20260629']

gd1 = defaultdict(lambda: defaultdict(int))
for i, row in enumerate(ws1.iter_rows(values_only=True)):
    if i == 0: continue
    v = list(row)
    season = str(v[10] or '').strip()
    if season not in seasons: continue
    q = v[23] or 0
    if not q: continue
    brand = str(v[3] or '').strip()
    gd = str(v[8] or '').strip()  # I열 = Grouped Div (index 8)
    if 'GAP' in brand.upper() or 'IPSS' in brand.upper():
        gd1[brand][gd] += q

wb1.close()
for brand in sorted(gd1.keys()):
    for gd, q in sorted(gd1[brand].items(), key=lambda x: -x[1]):
        print(f'  Brand={brand:20s} -> Grouped Div={gd:20s} qty={int(q):>10,}')
