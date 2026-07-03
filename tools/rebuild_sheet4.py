#!/usr/bin/env python3
"""Rebuild Sheet 4: Season > Factory rows vs Dept-Brand columns"""
import openpyxl, os, glob
from openpyxl.styles import Font, Border, Side, Alignment
from collections import defaultdict
from datetime import datetime

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
existing = os.path.join(BASE, '영업DATA_5Sheet_20260630.xlsx')
OUT = os.path.join(BASE, f'영업DATA_5Sheet_v2_{datetime.now().strftime("%Y%m%d")}.xlsx')

files = [f for f in glob.glob(os.path.join(BASE, '2)*')) if not os.path.basename(f).startswith('~$')]
F2 = max(files, key=os.path.getmtime)
print(f'Reading: {F2}')

wb2 = openpyxl.load_workbook(F2, read_only=True, data_only=True)
ws = wb2['Raw_20260608']

seasons = ['26 SPR', '26 SUM', '26 FAL', '26 HOL']
S4 = defaultdict(lambda: 0)
cnt = 0

for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    v = list(row); cnt += 1
    q = v[26] or 0
    season = str(v[12] or '').strip()
    if not (q and season in seasons): continue
    dept = str(v[5] or '').strip()
    brand = str(v[4] or '').strip()
    factory = str(v[2] or '').strip()
    S4[(dept, brand, factory, season)] += q

wb2.close()
print(f'Aggregated: {cnt} rows')

def dept_group(d):
    if d.startswith('BU1'): return 'BU1-영업1부'
    if d.startswith('BU2'):
        return 'BU2-영업1부' if '_D1' in d else 'BU2-영업2부'
    if d.startswith('BU3'):
        if '_D1' in d: return 'BU3-영업1부'
        if '_D2' in d: return 'BU3-영업2부'
        if '_D3' in d: return 'BU3-영업3부'
        if '_D4' in d: return 'BU3-영업4부'
    return 'BU4부'

DEPT_ORDER = ['BU1-영업1부','BU2-영업1부','BU2-영업2부',
              'BU3-영업1부','BU3-영업2부','BU3-영업3부','BU3-영업4부','BU4부']

# Build col structure: dept_group -> brand
dept_brands = defaultdict(set)
for (d, b, f, s), q in S4.items():
    if q: dept_brands[dept_group(d)].add(b)

db_pairs = []
for dg in DEPT_ORDER:
    for b in sorted(dept_brands.get(dg, set())):
        db_pairs.append((dg, b))

# Load existing workbook, replace sheet 4
wb = openpyxl.load_workbook(existing)
if '4.시즌부서공장' in wb.sheetnames:
    del wb['4.시즌부서공장']

ws4 = wb.create_sheet('4.시즌부서공장')
thin = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
hdr = Font(bold=True, size=10, name='Calibri')

ws4.cell(row=1, column=1, value='2026 Season 시즌별 Factory x 부서-브랜드 배정 수량').font = Font(bold=True, size=13, name='Calibri')
ws4.merge_cells(start_row=1, start_column=1, end_row=1, end_column=60)

# Row 3: Dept headers
ws4.cell(row=3, column=1, value='시즌').font = hdr
ws4.cell(row=3, column=2, value='Factory').font = hdr
col = 3
pair_col = {}
for dg in DEPT_ORDER:
    pairs = [(g,b) for (g,b) in db_pairs if g == dg]
    if not pairs: continue
    cs = col
    ws4.cell(row=3, column=col, value=dg).font = hdr
    ws4.cell(row=3, column=col).alignment = Alignment(horizontal='center')
    for g, b in pairs:
        pair_col[(g,b)] = col
        col += 1
    ce = col - 1
    if ce > cs:
        ws4.merge_cells(start_row=3, start_column=cs, end_row=3, end_column=ce)

# Row 4: Brand headers
for (dg, b), c in pair_col.items():
    cell = ws4.cell(row=4, column=c, value=b)
    cell.font = Font(bold=True, size=7, name='Calibri')
    cell.alignment = Alignment(horizontal='center', wrap_text=True, text_rotation=90)
    cell.border = thin

# Data: season > factory
r = 5
for season in seasons:
    factories = set()
    for (d, b, f, s), q in S4.items():
        if s == season and q > 0:
            factories.add(f)
    factories = sorted(factories)
    if not factories: continue
    
    start_r = r
    for fi, factory in enumerate(factories):
        ws4.cell(row=r+fi, column=2, value=factory).font = Font(size=9, name='Calibri')
        ws4.cell(row=r+fi, column=2).border = thin
        ws4.cell(row=r+fi, column=1).border = thin
        
        for (dg, b), c in pair_col.items():
            tot = 0
            for (d2, b2, f2, s2), q2 in S4.items():
                if b2 == b and f2 == factory and s2 == season and dept_group(d2) == dg:
                    tot += q2
            cell = ws4.cell(row=r+fi, column=c, value=int(tot) if tot else 0)
            cell.font = Font(size=9, name='Calibri'); cell.border = thin
            cell.alignment = Alignment(horizontal='right')
    
    # Season label (merged)
    ws4.merge_cells(start_row=start_r, start_column=1, end_row=start_r+len(factories)-1, end_column=1)
    cell = ws4.cell(row=start_r, column=1, value=season)
    cell.font = Font(bold=True, size=10, name='Calibri')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    
    r += len(factories) + 1  # blank separator row

ws4.column_dimensions['A'].width = 12
ws4.column_dimensions['B'].width = 32

wb.save(OUT)
print(f'Saved: {OUT}')
import os; print(f'Size: {os.path.getsize(OUT)/1024:.0f} KB')
