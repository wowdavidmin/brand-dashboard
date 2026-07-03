#!/usr/bin/env python3
"""Deep comparison: reference BU1/BU2 brand structure vs file #2 brands"""
import openpyxl
from collections import defaultdict

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'

# 1. Read reference old report - Sheet 1 structure
print('=== Reference: Sheet 1 brand structure ===')
wb = openpyxl.load_workbook(BASE + r'\참고_★ 2021~2026 영업 DATA FILE_전달용_v1.2_20260609의 백업.xlsx', read_only=True, data_only=True)
ws = wb['1.21~26시즌현황']

# Check headers
r = 0
dept_headers = {}
brand_headers = {}
for row in ws.iter_rows(values_only=True):
    r += 1
    vals = list(row)
    if r == 5:  # 부서 row
        for c, v in enumerate(vals):
            if v and str(v).strip():
                dept_headers[c] = str(v).strip()
    if r == 6:  # BUYER row
        for c, v in enumerate(vals):
            if v and str(v).strip():
                brand_headers[c] = str(v).strip()
    if r >= 12:  # Only first 2 sections (ORDER qty and Total sales)
        break

print('BUYER columns in reference sheet 1:')
buyer_list = []
for c in sorted(brand_headers.keys()):
    buyer = brand_headers[c]
    dept = dept_headers.get(c, '')
    buyer_list.append((buyer, c, dept))
    print(f'  col {c:3d}: {dept:15s} | {buyer:25s}')

print(f'\nTotal buyers: {len(buyer_list)}')
wb.close()

# 2. Read file #2 to see all brand names used
print('\n=== File #2: All brand names ===')
f2 = BASE + r'\2) 영업DATA_Raw_v1.3_20260608_openclaw 작업용.xlsx'
wb2 = openpyxl.load_workbook(f2, read_only=True, data_only=True)
ws2 = wb2['Raw_20260608']

f2_brands = defaultdict(lambda: [0,0,0,0,0,0])  # brand -> [qty, amt, net, dc, profit, em]
seasons = {'26 SPR','26 SUM','26 FAL','26 HOL'}
cnt = 0

for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0: continue
    v = list(row); cnt += 1
    q = v[26] or 0; a = v[27] or 0; n = v[29] or 0
    dc = v[48] or 0; prof = v[50] or 0; em = v[56] or 0
    brand = str(v[4] or '').strip()
    season = str(v[12] or '').strip()
    if not (q and a and season in seasons): continue
    f2_brands[brand] = [f2_brands[brand][i] + x for i, x in enumerate([q, a, n, dc, prof, em])]

wb2.close()
print(f'Total rows: {cnt}')
print(f'Total brands: {len(f2_brands)}')

# Show BU1 and BU2 related brands
print('\n--- GAP-related brands in file #2 ---')
for brand, d in sorted(f2_brands.items()):
    if 'GAP' in brand.upper() or 'GAP' in brand or 'BRFS' in brand or 'MOTT' in brand or 'BOW' in brand or 'MADHAPPY' in brand:
        qty = int(d[0])
        amt = int(d[1])
        up = round(d[1]/d[0], 2) if d[0] else 0
        print(f'  {brand:30s} qty={qty:>10,}  amt=${amt:>10,}  unit=${up}')

print('\n--- OLD NAVY / ON brands in file #2 ---')
for brand, d in sorted(f2_brands.items()):
    if 'OLD NAVY' in brand.upper() or brand.upper().startswith('ON ') or 'ON MEN' in brand.upper() or 'ON WOM' in brand.upper() or 'ON KID' in brand.upper():
        qty = int(d[0])
        amt = int(d[1])
        up = round(d[1]/d[0], 2) if d[0] else 0
        print(f'  {brand:30s} qty={qty:>10,}  amt=${amt:>10,}  unit=${up}')

print('\n--- WALMART / QUINCE brands in file #2 ---')
for brand, d in sorted(f2_brands.items()):
    if 'WALMART' in brand.upper() or 'WAL-MART' in brand or 'QUINCE' in brand.upper():
        qty = int(d[0])
        amt = int(d[1])
        up = round(d[1]/d[0], 2) if d[0] else 0
        print(f'  {brand:30s} qty={qty:>10,}  amt=${amt:>10,}  unit=${up}')

# Check for brands in reference that might be missing in file #2
print('\n=== Brands in reference but potentially missing/misaligned in file #2 ===')
ref_brands_BU1_BU2 = {
    'GAP KIDS': 7730913, "GAP MEN'S": 6271624, "GAP WMN'S": 8806969,
    'IPSS GAP': 2202054, 'BRFS': 3236380,
    'OLD NAVY': None, "ON  MEN'S": None, "ON  WMN'S": None, "ON KID'S & BABY": None,
    'WALMART CANADA': 4160576, "WALMART D29 (INTIMATE)": 11522672,
    "WALMART D26 (BABY)": 15038858, 'QUINCE': 1581884,
}
for ref_brand in sorted(ref_brands_BU1_BU2.keys()):
    # Find closest match in f2_brands
    matches = [(b, d[0]) for b, d in f2_brands.items() 
               if ref_brand[:6].upper() in b.upper() or b[:6].upper() in ref_brand.upper()]
    if not matches:
        print(f'  MISSING: {ref_brand:30s} -> no match found in file #2')
    else:
        for b, qty in matches:
            print(f'  {ref_brand:30s} -> {b:25s} (qty={int(qty):>10,})')
