#!/usr/bin/env python3
"""Rebuild Sheet 4 from v3, save as v4"""
import openpyxl, os, shutil
from openpyxl.styles import Font, Border, Side, Alignment
from collections import defaultdict

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
src = os.path.join(BASE, '리포트_5Sheet_최종_v3_20260630.xlsx')
out = os.path.join(BASE, '리포트_5Sheet_최종_v4_20260630.xlsx')
shutil.copy2(src, out)

wb = openpyxl.load_workbook(out)
ws4 = wb['4.시즌부서공장']
rows = list(ws4.iter_rows(values_only=True))

# Row[2] = brand names, Row[1] = dept names
brand_row = rows[2]
brand_to_dept = {}
for c in range(2, len(brand_row)):
    bn = str(brand_row[c] or '').strip()
    dn = str(rows[1][c] or '').strip() if c < len(rows[1]) else ''
    if bn and bn not in ('', 'Factory'):
        brand_to_dept[bn] = dn

DEPT_ORDER = ['BU1-영업1부','BU2-영업1부','BU2-영업2부','BU3-영업1부','BU3-영업2부','BU3-영업3부','BU3-영업4부','BU4부','BU5-영업5부','GTBU-과테말라']

# Parse data
data = []
cur_season = ''
for ri in range(3, len(rows)):
    row = rows[ri]
    s = str(row[0] or '').strip()
    if s in ('26 SPR','26 SUM','26 FAL','26 HOL'):
        cur_season = s
    fac = str(row[1] or '').strip()
    if not fac or not cur_season: continue
    for c in range(2, len(row)):
        bn = str(brand_row[c] or '').strip()
        if not bn or bn in ('', 'Factory'): continue
        q = row[c] if row[c] else 0
        try: q = int(q)
        except: q = 0
        if q > 0:
            data.append((bn, fac, cur_season, q))

# Deduplicate
agg = defaultdict(int)
for bn, fac, s, q in data:
    agg[(bn, fac, s)] += q

del wb['4.시즌부서공장']
ws4n = wb.create_sheet('4.시즌부서공장')
thin = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
hdr = Font(bold=True, size=10, name='Calibri')

ws4n.cell(row=1, column=1, value='2026 Season 브랜드/Factory/시즌별 부서 배정 수량').font = Font(bold=True, size=13, name='Calibri')
ws4n.merge_cells('A1:ZZ1')
ws4n.cell(row=3, column=1, value='브랜드').font = hdr
ws4n.cell(row=3, column=2, value='Factory').font = hdr
ws4n.cell(row=3, column=3, value='시즌').font = hdr
col = 4; dc = {}
for dn in DEPT_ORDER:
    ws4n.cell(row=3, column=col, value=dn).font = hdr; dc[dn] = col; col += 1

FMT = '#,##0'
r = 4
for (bn, fac, s), qty in sorted(agg.items()):
    ws4n.cell(row=r, column=1, value=bn)
    ws4n.cell(row=r, column=2, value=fac)
    ws4n.cell(row=r, column=3, value=s)
    for ci in range(1,4):
        ws4n.cell(row=r, column=ci).font = Font(size=10); ws4n.cell(row=r, column=ci).border = thin
    dn = brand_to_dept.get(bn, 'BU4부')
    for d in DEPT_ORDER:
        c = dc[d]; v = qty if d == dn else 0
        cell = ws4n.cell(row=r, column=c, value=v if v else 0)
        cell.number_format = FMT; cell.font = Font(size=10); cell.border = thin
        cell.alignment = Alignment(horizontal='right')
    r += 1

order = ['1.시즌현황','2.시즌별','3.월별','4.시즌부서공장','5.시즌공장']
for i, name in enumerate(order):
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        if idx != i: wb.move_sheet(name, offset=i-idx)

print(f'Rows: {len(agg)}')
wb.save(out)
import os; print(f'Saved: {out} ({os.path.getsize(out)/1024:.0f} KB)')
wb.close()
