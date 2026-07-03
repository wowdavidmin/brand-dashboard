#!/usr/bin/env python3
import openpyxl
from collections import defaultdict

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'

# Check reference for GAP
print('=== Reference Sheet 2: GAP brands ===')
wb = openpyxl.load_workbook(BASE + r'\참고_★ 2021~2026 영업 DATA FILE_전달용_v1.2_20260609의 백업.xlsx', read_only=True, data_only=True)
ws = wb['2.21~26시즌별']
r = 0
for row in ws.iter_rows(values_only=True):
    r += 1
    if r < 7: continue
    vals = list(row)
    name = str(vals[0] or '').strip()
    if not name or '요약' in name: continue
    if name.startswith('[') or name.startswith('*'): continue
    if name in ('26 SPR','26 SUM','26 FAL','26 HOL',''): continue
    try:
        total = int(float(vals[13])) if vals[13] else 0
        if 'GAP' in name.upper():
            print(f'  {name:25s} TOTAL={total:>10,}')
    except: pass
wb.close()

# Check file #2 for GAP brands and buyers
print('\n=== File #2: All data with GAP in buyer or brand ===')
f2 = BASE + r'\2) 영업DATA_Raw_v1.3_20260608_openclaw 작업용.xlsx'
wb2 = openpyxl.load_workbook(f2, read_only=True, data_only=True)
ws2 = wb2['Raw_20260608']

info = defaultdict(lambda: [0, 0])
seasons = {'26 SPR','26 SUM','26 FAL','26 HOL'}
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0: continue
    v = list(row)
    q = v[26] or 0
    a = v[27] or 0
    season = str(v[12] or '').strip()
    if not (q and season in seasons): continue
    buyer = str(v[3] or '').strip()
    brand = str(v[4] or '').strip()
    if 'GAP' in buyer.upper() or 'GAP' in brand.upper():
        key = f'{buyer} / {brand}'
        info[key][0] += q
        info[key][1] += a

wb2.close()
for key in sorted(info.keys()):
    q, a = info[key]
    up = a/q if q else 0
    print(f'  {key:50s} qty={int(q):>10,}  amt=${int(a):>10,}  unit=${up:.2f}')
