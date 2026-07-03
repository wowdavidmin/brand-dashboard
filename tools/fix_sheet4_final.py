#!/usr/bin/env python3
"""Rebuild Sheet 4: dept|brand|factory|seasons(qty+%)|total"""
import openpyxl, os, shutil
from openpyxl.styles import Font, Border, Side, Alignment
from collections import defaultdict

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
src = os.path.join(BASE, '리포트_5Sheet_최종_v3_20260630.xlsx')
out = os.path.join(BASE, '리포트_5Sheet_최종_v8_20260630.xlsx')
shutil.copy2(src, out)

wb = openpyxl.load_workbook(out)
ws4 = wb['4.시즌부서공장']
rows = list(ws4.iter_rows(values_only=True))

# Parse brand->dept mapping
brand_to_dept = {}
last_dept = ''
for c in range(2, len(rows[2])):
    bn = str(rows[2][c] or '').strip()
    dn = str(rows[1][c] or '').strip() if c < len(rows[1]) else ''
    if dn: last_dept = dn
    if bn and bn not in ('', 'Factory'):
        brand_to_dept[bn] = last_dept

# Parse: (dept, brand, factory, season) -> qty
data = defaultdict(int)
cur_season = ''
for ri in range(3, len(rows)):
    row = rows[ri]
    s = str(row[0] or '').strip()
    if s in ('26 SPR','26 SUM','26 FAL','26 HOL'): cur_season = s
    fac = str(row[1] or '').strip()
    if not fac or not cur_season: continue
    for c in range(2, len(row)):
        bn = str(rows[2][c] or '').strip()
        if not bn or bn in ('', 'Factory'): continue
        q = row[c] if row[c] else 0
        try: q = int(float(str(q)))
        except: q = 0
        if q > 0:
            dn = brand_to_dept.get(bn, 'BU4부')
            data[(dn, bn, fac, cur_season)] += q

del wb['4.시즌부서공장']
ws4n = wb.create_sheet('4.시즌부서공장')

thin = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
hdr = Font(bold=True, size=10, name='Calibri')

ws4n.cell(row=1, column=1, value='2026 Season 조직-브랜드-Factory별 시즌 생산 수량 및 비중').font = Font(bold=True, size=13, name='Calibri')
ws4n.merge_cells('A1:Z1')

FMT = '#,##0'
seasons_list = ['26 SPR', '26 SUM', '26 FAL', '26 HOL']

# Row 3: season headers (merged across qty+%)
col = 4
se_col = {}
for s in seasons_list:
    ws4n.cell(row=3, column=col, value=s).font = hdr
    ws4n.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+1)
    ws4n.cell(row=4, column=col, value='수량').font = hdr
    ws4n.cell(row=4, column=col+1, value='비중%').font = hdr
    se_col[s] = (col, col+1)
    col += 2

ws4n.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col)
ws4n.cell(row=3, column=col, value='Total').font = hdr
ws4n.cell(row=4, column=col, value='수량').font = hdr
total_col = col

ws4n.cell(row=3, column=1, value='조직').font = hdr
ws4n.cell(row=3, column=2, value='브랜드').font = hdr
ws4n.cell(row=3, column=3, value='Factory').font = hdr

for ci in [1,2,3]:
    ws4n.cell(row=4, column=ci).font = hdr

# Group data by (dept, brand, factory)
grouped = defaultdict(lambda: defaultdict(int))
for (dn, bn, fac, s), q in data.items():
    grouped[(dn, bn, fac)][s] += q

r = 5
last_dn = None
for (dn, bn, fac), sq in sorted(grouped.items()):
    total = sum(sq.values())
    
    if dn != last_dn:
        ws4n.cell(row=r, column=1, value=dn)
        last_dn = dn
    ws4n.cell(row=r, column=2, value=bn)
    ws4n.cell(row=r, column=3, value=fac)
    for ci in range(1, 4):
        ws4n.cell(row=r, column=ci).font = Font(size=10)
        ws4n.cell(row=r, column=ci).border = thin
        ws4n.cell(row=r, column=ci).alignment = Alignment(horizontal='left')
    
    for s in seasons_list:
        cq, cp = se_col[s]
        qty = sq.get(s, 0)
        pct = round(qty / total * 100, 1) if total else 0
        
        cell_q = ws4n.cell(row=r, column=cq, value=qty if qty else 0)
        cell_q.number_format = FMT; cell_q.font = Font(size=10); cell_q.border = thin
        cell_q.alignment = Alignment(horizontal='right')
        
        cell_p = ws4n.cell(row=r, column=cp, value=pct if qty else 0)
        cell_p.number_format = '0.0'; cell_p.font = Font(size=10); cell_p.border = thin
        cell_p.alignment = Alignment(horizontal='right')
    
    cell_t = ws4n.cell(row=r, column=total_col, value=total if total else 0)
    cell_t.number_format = FMT; cell_t.font = Font(size=10); cell_t.border = thin
    cell_t.alignment = Alignment(horizontal='right')
    
    r += 1

# Column widths
ws4n.column_dimensions['A'].width = 16
ws4n.column_dimensions['B'].width = 22
ws4n.column_dimensions['C'].width = 32

# Fix order
order = ['1.시즌현황','2.시즌별','3.월별','4.시즌부서공장','5.시즌공장']
for i, name in enumerate(order):
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        if idx != i: wb.move_sheet(name, offset=i-idx)

print(f'Rows: {len(grouped)}')
wb.save(out)
import os; print(f'Saved: {out} ({os.path.getsize(out)/1024:.0f} KB)')
wb.close()
