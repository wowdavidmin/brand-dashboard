#!/usr/bin/env python3
"""
약진통상 영업데이터 리포트 자동 생성기
사용법: python3 weekly_report.py [raw_file] [output_name]
"""
import sys, os, pickle
from collections import defaultdict
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Border, Side

BASE_DIR = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'

def find_latest_file(prefix):
    """Find latest file matching prefix in BASE_DIR"""
    import glob
    files = glob.glob(os.path.join(BASE_DIR, f'*{prefix}*'))
    if not files:
        return None
    return max(files, key=os.path.getmtime)

def aggregate_raw_data(f2):
    """Read raw data and build all aggregations"""
    print(f'Reading: {f2}')
    wb2 = openpyxl.load_workbook(f2, read_only=True, data_only=True)
    ws = wb2['Raw_20260608']

    S = {
        '시즌_data': defaultdict(lambda: [0, 0, 0]),
        '시즌_바이어': set(),
        '브랜드시즌': defaultdict(lambda: [0, 0, 0]),
        '브랜드셋': set(),
        '월별브랜드': defaultdict(lambda: [0, 0, 0]),
        '부서공장시즌': defaultdict(lambda: 0),
    }
    시즌 = {'26 SPR', '26 SUM', '26 FAL', '26 HOL'}
    row_cnt = 0

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0: continue
        v = list(row)
        row_cnt += 1

        qty = v[26] or 0
        amt = v[27] or 0
        net = v[29] or 0
        buyer = str(v[3] or '').strip()
        brand = str(v[4] or '').strip()
        season = str(v[12] or '').strip()
        dept = str(v[5] or '').strip()
        factory = str(v[2] or '').strip()
        month = str(v[21] or '').strip()

        if not (qty and amt and season in 시즌): continue

        S['시즌_data'][(season, buyer)] = [S['시즌_data'][(season, buyer)][i] + v for i, v in enumerate([qty, amt, net])]
        S['시즌_바이어'].add(buyer)
        S['브랜드시즌'][(brand, season)] = [S['브랜드시즌'][(brand, season)][i] + v for i, v in enumerate([qty, amt, net])]
        S['브랜드셋'].add(brand)
        if month:
            S['월별브랜드'][(brand, month)] = [S['월별브랜드'][(brand, month)][i] + v for i, v in enumerate([qty, amt, net])]
        S['부서공장시즌'][(dept, brand, factory, season)] += qty
        if row_cnt % 30000 == 0: print(f'  {row_cnt}행...')

    wb2.close()
    print(f'완료: {row_cnt}행')
    return S


def guess_dept(buyer):
    b = buyer.upper()
    if 'GAP' in b and 'IPSS' not in b: return 'BU1-영업1부'
    if 'IPSS GAP' in b: return 'BU1-영업1부'
    if b in ('BRFS',): return 'BU1-영업1부'
    if b in ('MADHAPPY',): return 'BU1-영업1부'
    if 'MOTT' in b or 'BOW' in b: return 'BU1-영업1부'
    if 'OLD NAVY' in b: return 'BU2-영업1부'
    if b.startswith('ON '): return 'BU2-영업1부'
    if 'WALMART' in b: return 'BU2-영업2부'
    if 'QUINCE' in b: return 'BU2-영업2부'
    if b in ('BR',): return 'BU3-영업1부'
    if 'IPSS ON' in b: return 'BU3-영업1부'
    if 'KOHL' in b: return 'BU3-영업2부'
    if 'URBAN' in b: return 'BU3-영업2부'
    if 'AERIE' in b: return 'BU3-영업3부'
    if 'OFFLINE' in b: return 'BU3-영업3부'
    if 'MLB' in b: return 'BU3-영업4부'
    return 'BU4부'


def guess_brand_dept(brand):
    b = brand.upper()
    if 'GAP' in b and 'IPSS' not in b: return 'BU1-영업1부'
    if 'IPSS GAP' in b: return 'BU1-영업1부'
    if b in ('BRFS',): return 'BU1-영업1부'
    if b in ('MADHAPPY',): return 'BU1-영업1부'
    if 'MOTT' in b or 'BOW' in b: return 'BU1-영업1부'
    if 'OLD NAVY' in b: return 'BU2-영업1부'
    if b.startswith('ON '): return 'BU2-영업1부'
    if 'WALMART' in b: return 'BU2-영업2부'
    if 'QUINCE' in b: return 'BU2-영업2부'
    if b in ('BR',): return 'BU3-영업1부'
    if 'IPSS ON' in b: return 'BU3-영업1부'
    if 'KOHL' in b: return 'BU3-영업2부'
    if 'URBAN' in b: return 'BU3-영업2부'
    if 'AERIE' in b: return 'BU3-영업3부'
    if 'OFFLINE' in b: return 'BU3-영업3부'
    if 'MLB' in b: return 'BU3-영업4부'
    return 'BU4부'


def make_report(S, output_path):
    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    hdr = Font(bold=True, size=11)
    ttl = Font(bold=True, size=13)
    seasons = ['26 SPR', '26 SUM', '26 FAL', '26 HOL']
    months = [f'{m}월' for m in range(1, 13)]
    dept_order = ['BU1-영업1부', 'BU2-영업1부', 'BU2-영업2부', 'BU3-영업1부', 'BU3-영업2부', 'BU3-영업3부', 'BU3-영업4부', 'BU4부']

    # Group brands by dept
    dept_brands = defaultdict(set)
    for brand in S['브랜드셋']:
        dept_brands[guess_brand_dept(brand)].add(brand)

    wb = openpyxl.Workbook()

    # ===== SHEET 1: 시즌현황 =====
    ws1 = wb.active
    ws1.title = '1.시즌현황'
    ws1.cell(row=1, column=1, value='[약진통상] 2026 Season 부서별 브랜드별 - 수량/매출/마진').font = ttl
    ws1.merge_cells('A1:F1')
    ws1.cell(row=2, column=1, value='데이터기준: ' + datetime.now().strftime('%Y.%m.%d')).font = Font(size=9, italic=True)
    ws1.merge_cells('A2:F2')

    for ci, h in enumerate(['시즌', '부서', '바이어', 'ORDER수량', '총매출액', '순매출액'], 1):
        c = ws1.cell(row=4, column=ci, value=h); c.font = hdr; c.border = thin

    r = 5
    for season in seasons:
        items = [(b, S['시즌_data'].get((season, b), [0, 0, 0])) for b in S['시즌_바이어']]
        items = [(b, d) for b, d in items if d[0] > 0]
        items.sort(key=lambda x: -x[1][0])
        for buyer, d in items:
            for ci, v in enumerate([season, guess_dept(buyer), buyer, int(d[0]), round(d[1], 2), round(d[2], 2)], 1):
                c = ws1.cell(row=r, column=ci, value=v); c.border = thin
            r += 1
        sq = sum(S['시즌_data'].get((season, b), [0, 0, 0])[0] for b in S['시즌_바이어'])
        sa = sum(S['시즌_data'].get((season, b), [0, 0, 0])[1] for b in S['시즌_바이어'])
        sn = sum(S['시즌_data'].get((season, b), [0, 0, 0])[2] for b in S['시즌_바이어'])
        for ci, v in enumerate([season, '소계', '', int(sq), round(sa, 2), round(sn, 2)], 1):
            c = ws1.cell(row=r, column=ci, value=v); c.font = Font(bold=True); c.border = thin
        r += 1

    tq = sum(S['시즌_data'][(s, b)][0] for s in seasons for b in S['시즌_바이어'])
    ta = sum(S['시즌_data'][(s, b)][1] for s in seasons for b in S['시즌_바이어'])
    tn = sum(S['시즌_data'][(s, b)][2] for s in seasons for b in S['시즌_바이어'])
    for ci, v in enumerate(['총합계', '', '', int(tq), round(ta, 2), round(tn, 2)], 1):
        c = ws1.cell(row=r, column=ci, value=v); c.font = Font(bold=True, size=12); c.border = thin
    ws1.column_dimensions['A'].width = 12; ws1.column_dimensions['B'].width = 14
    ws1.column_dimensions['C'].width = 22; ws1.column_dimensions['D'].width = 14
    ws1.column_dimensions['E'].width = 16; ws1.column_dimensions['F'].width = 16
    print('Sheet 1 done')

    # ===== SHEET 2: 시즌별 =====
    ws2 = wb.create_sheet('2.시즌별')
    ws2.cell(row=1, column=1, value='[약진통상] 2026 Season별 부서별 브랜드디비전별 매출수량/매출액').font = ttl
    ws2.merge_cells('A1:P1')
    ws2.cell(row=3, column=1, value='브랜드/부서').font = hdr

    col2 = {}
    c = 2
    for s in seasons:
        ws2.cell(row=3, column=c, value=s).font = hdr
        ws2.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c+2)
        ws2.cell(row=4, column=c, value='수량').font = hdr
        ws2.cell(row=4, column=c+1, value='총매출액').font = hdr
        ws2.cell(row=4, column=c+2, value='순매출액').font = hdr
        col2[s] = (c, c+1, c+2); c += 3
    ws2.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c+3)
    ws2.cell(row=3, column=c, value='TOTAL').font = hdr
    ws2.cell(row=4, column=c, value='수량').font = hdr
    ws2.cell(row=4, column=c+1, value='총매출액').font = hdr
    ws2.cell(row=4, column=c+2, value='순매출액').font = hdr
    ws2.cell(row=4, column=c+3, value='마진율%').font = hdr
    tc2 = c

    r = 5
    for dn in dept_order:
        brands = sorted(dept_brands.get(dn, set()))
        if not brands: continue
        for brand in brands:
            ws2.cell(row=r, column=1, value=brand)
            for s in seasons:
                d = S['브랜드시즌'].get((brand, s), [0, 0, 0])
                cq, ca, cn_ = col2[s]
                ws2.cell(row=r, column=cq, value=int(d[0])); ws2.cell(row=r, column=ca, value=round(d[1], 2)); ws2.cell(row=r, column=cn_, value=round(d[2], 2))
            tq = sum(S['브랜드시즌'].get((brand, s), [0, 0, 0])[0] for s in seasons)
            ta = sum(S['브랜드시즌'].get((brand, s), [0, 0, 0])[1] for s in seasons)
            tn = sum(S['브랜드시즌'].get((brand, s), [0, 0, 0])[2] for s in seasons)
            ws2.cell(row=r, column=tc2, value=int(tq)); ws2.cell(row=r, column=tc2+1, value=round(ta, 2)); ws2.cell(row=r, column=tc2+2, value=round(tn, 2))
            ws2.cell(row=r, column=tc2+3, value=round((ta-tn)/ta*100, 2) if ta else 0)
            for ci in range(1, tc2+4): ws2.cell(row=r, column=ci).border = thin
            r += 1
        # subtotal
        ws2.cell(row=r, column=1, value=f'{dn} 요약').font = Font(bold=True)
        for s in seasons:
            dq = sum(S['브랜드시즌'].get((b, s), [0, 0, 0])[0] for b in brands)
            da = sum(S['브랜드시즌'].get((b, s), [0, 0, 0])[1] for b in brands)
            dn_ = sum(S['브랜드시즌'].get((b, s), [0, 0, 0])[2] for b in brands)
            cq, ca, cn_ = col2[s]
            ws2.cell(row=r, column=cq, value=int(dq)).font = Font(bold=True)
            ws2.cell(row=r, column=ca, value=round(da, 2)).font = Font(bold=True)
            ws2.cell(row=r, column=cn_, value=round(dn_, 2)).font = Font(bold=True)
        tq = sum(sum(S['브랜드시즌'].get((b, s), [0, 0, 0])[0] for s in seasons) for b in brands)
        ta = sum(sum(S['브랜드시즌'].get((b, s), [0, 0, 0])[1] for s in seasons) for b in brands)
        tn = sum(sum(S['브랜드시즌'].get((b, s), [0, 0, 0])[2] for s in seasons) for b in brands)
        ws2.cell(row=r, column=tc2, value=int(tq)).font = Font(bold=True)
        ws2.cell(row=r, column=tc2+1, value=round(ta, 2)).font = Font(bold=True)
        ws2.cell(row=r, column=tc2+2, value=round(tn, 2)).font = Font(bold=True)
        ws2.cell(row=r, column=tc2+3, value=round((ta-tn)/ta*100, 2) if ta else 0).font = Font(bold=True)
        for ci in range(1, tc2+4): ws2.cell(row=r, column=ci).border = thin
        r += 1
    ws2.column_dimensions['A'].width = 24
    print('Sheet 2 done')

    # ===== SHEET 3: 월별 =====
    ws3 = wb.create_sheet('3.월별')
    ws3.cell(row=1, column=1, value='[약진통상] 2026년 월별 부서별 브랜드디비전별 매출수량/매출액').font = ttl
    ws3.merge_cells('A1:AH1')
    ws3.cell(row=3, column=1, value='브랜드/부서').font = hdr

    col3 = {}
    c = 2
    for mn in months:
        ws3.cell(row=3, column=c, value=mn).font = hdr
        ws3.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c+2)
        ws3.cell(row=4, column=c, value='수량').font = hdr
        ws3.cell(row=4, column=c+1, value='총매출액').font = hdr
        ws3.cell(row=4, column=c+2, value='순매출액').font = hdr
        col3[mn] = (c, c+1, c+2); c += 3

    r = 5
    for dn in dept_order:
        brands = sorted(dept_brands.get(dn, set()))
        if not brands: continue
        for brand in brands:
            ws3.cell(row=r, column=1, value=brand)
            for mn in months:
                d = S['월별브랜드'].get((brand, mn), [0, 0, 0])
                cq, ca, cn_ = col3[mn]
                ws3.cell(row=r, column=cq, value=int(d[0]))
                ws3.cell(row=r, column=ca, value=round(d[1], 2))
                ws3.cell(row=r, column=cn_, value=round(d[2], 2))
            for ci in range(1, c): ws3.cell(row=r, column=ci).border = thin
            r += 1
        ws3.cell(row=r, column=1, value=f'{dn} 요약').font = Font(bold=True)
        for mn in months:
            dq = sum(S['월별브랜드'].get((b, mn), [0, 0, 0])[0] for b in brands)
            da = sum(S['월별브랜드'].get((b, mn), [0, 0, 0])[1] for b in brands)
            dn_ = sum(S['월별브랜드'].get((b, mn), [0, 0, 0])[2] for b in brands)
            cq, ca, cn_ = col3[mn]
            ws3.cell(row=r, column=cq, value=int(dq)).font = Font(bold=True)
            ws3.cell(row=r, column=ca, value=round(da, 2)).font = Font(bold=True)
            ws3.cell(row=r, column=cn_, value=round(dn_, 2)).font = Font(bold=True)
        for ci in range(1, c): ws3.cell(row=r, column=ci).border = thin
        r += 1
    print('Sheet 3 done')

    # ===== SHEET 4: 시즌부서공장 =====
    ws4 = wb.create_sheet('4.시즌부서공장')
    ws4.cell(row=1, column=1, value='[약진통상] 2026 Season 부서별 브랜드별 공장별 오더 수량').font = ttl
    ws4.merge_cells('A1:K1')
    for ci, h in enumerate(['부서', '브랜드', 'Factory'], 1):
        ws4.cell(row=3, column=ci, value=h).font = hdr
    col4 = {}
    c = 4
    for s in seasons:
        ws4.cell(row=3, column=c, value=s).font = hdr
        ws4.cell(row=3, column=c+1, value='비중').font = hdr
        col4[s] = c; c += 2
    ws4.cell(row=3, column=c, value='Total').font = hdr
    ws4.cell(row=3, column=c+1, value='비중').font = hdr
    tc4 = c

    r = 4
    for dn in dept_order:
        prefix = dn.split('-')[0]
        items = [(k, v) for k, v in S['부서공장시즌'].items() if k[0].startswith(prefix)]
        if not items: continue
        dg = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
        for (dp, brand, factory, season), qty in items:
            dg[brand][factory][season] += qty
        for brand in sorted(dg.keys()):
            factories = sorted(dg[brand].keys())
            for fi, factory in enumerate(factories):
                ws4.cell(row=r, column=1, value=dn if fi == 0 else '')
                ws4.cell(row=r, column=2, value=brand if fi == 0 else '')
                ws4.cell(row=r, column=3, value=factory)
                for s in seasons:
                    ws4.cell(row=r, column=col4[s], value=int(dg[brand][factory].get(s, 0)))
                t = sum(dg[brand][factory].get(s, 0) for s in seasons)
                ws4.cell(row=r, column=tc4, value=int(t))
                for ci in range(1, tc4+2): ws4.cell(row=r, column=ci).border = thin
                r += 1
    print('Sheet 4 done')

    wb.save(output_path)
    print(f'Saved: {output_path}')


if __name__ == '__main__':
    import glob
    if len(sys.argv) > 1:
        raw_file = sys.argv[1]
    else:
        files = sorted([f for f in glob.glob(os.path.join(BASE_DIR, '*영업DATA_Raw*')) if not os.path.basename(f).startswith('~$')], key=os.path.getmtime, reverse=True)
        raw_file = files[0] if files else None
        if not raw_file:
            files = sorted([f for f in glob.glob(os.path.join(BASE_DIR, '*2)*')) if not os.path.basename(f).startswith('~$')], key=os.path.getmtime, reverse=True)
            raw_file = files[0] if files else None

    if not raw_file:
        print('No raw data file found!')
        sys.exit(1)

    out_name = f'영업DATA_Report_{datetime.now().strftime("%Y%m%d")}.xlsx'
    if len(sys.argv) > 2:
        out_name = sys.argv[2]
    out_path = os.path.join(BASE_DIR, out_name)

    data = aggregate_raw_data(raw_file)
    make_report(data, out_path)
    print(f'\n완료! 파일: {out_path}')
