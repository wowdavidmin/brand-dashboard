#!/usr/bin/env python3
import openpyxl
BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
f = BASE + r'\영업DATA_5Sheet_FINAL_20260630.xlsx'
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb['1.시즌현황']

brands = {}
total_row = None
for row in ws.iter_rows(values_only=True):
    if row[0] == '시즌':
        for c, v in enumerate(row):
            if c >= 1 and v and isinstance(v, str) and v not in ('수량','총매출','순매출','평균단가','마진율%','CM%',''):
                brands[c] = v.strip()
    if row[0] == 'Total':
        total_row = row
        break

print('=== BU1 Detail ===')
bu1_list = ['GAP', 'IPSS GAP', 'BRFS', 'Madhappy', 'Mott&Bow', 'Gap OUTLET']
bu1_total = 0
for c in sorted(brands.keys()):
    b = brands[c]
    v = total_row[c] if total_row and total_row[c] else 0
    if not isinstance(v, (int, float)): v = 0
    print(f'  {b:20s}: {int(v):>10,}')
    if b in bu1_list:
        bu1_total += int(v)

dash = '-' * 32
print(f'  {dash}')
print(f'  {"BU1 합계":20s}: {bu1_total:>10,}')
print(f'  {"참고(2주전)":20s}: 28,287,329')
print(f'  {"증감":20s}: {bu1_total-28287329:>+10,}')
pct = (bu1_total - 28287329) / 28287329 * 100
print(f'  {"증감률":20s}: {pct:+.1f}%')

print()
gap_col = [c for c, b in brands.items() if b == 'GAP'][0]
gap_v = int(total_row[gap_col]) if total_row and total_row[gap_col] else 0
gap_outlet_col = [c for c, b in brands.items() if b == 'Gap OUTLET'][0]
gap_outlet_v = int(total_row[gap_outlet_col]) if total_row and total_row[gap_outlet_col] else 0
print(f'GAP: {gap_v:,}  (2주전 22,500,707)')
print(f'Gap OUTLET: {gap_outlet_v:,}  (신규 할인채널)')
print(f'GAP+OUTLET 합계: {gap_v + gap_outlet_v:,}')

print()
print(f'=== Sheet list ===')
for s in wb.sheetnames:
    ws2 = wb[s]
    cnt = sum(1 for _ in ws2.iter_rows(values_only=True))
    print(f'  [{s}] {cnt}행')

wb.close()
