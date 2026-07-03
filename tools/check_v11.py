#!/usr/bin/env python3
import openpyxl
f = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week\리포트_5Sheet_최종_v11_20260630.xlsx'
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb['1.시즌현황']
print('=== ORDER 수량 - Total row (col 1-13) ===')
for row in ws.iter_rows(values_only=True):
    if row[0] == 'Total':
        for c in range(1, 14):
            v = row[c]
            s = str(v)[:15] if v is not None else '(nil)'
            print(f'  col{c}: {s}')
        break
print()
print('=== Season rows with S.TOTAL ===')
for row in ws.iter_rows(values_only=True):
    if row[0] in ('26 SPR','26 SUM','26 FAL','26 HOL','Total'):
        g = row[1] if row[1] else '-'
        s = row[6] if len(row)>6 and row[6] else '-'
        print(f'  {str(row[0]):8s} GAP={str(g)[:10]:>10s}  S.TOTAL(BU1)={str(s)[:10]:>10s}')
wb.close()
