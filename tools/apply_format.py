#!/usr/bin/env python3
"""Apply formatting: 0->dash, comma separators, % for shares"""
import openpyxl, os, shutil

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
src = os.path.join(BASE, '리포트_5Sheet_최종_v8_20260630.xlsx')
out = os.path.join(BASE, '리포트_5Sheet_최종_FINAL_20260630.xlsx')
shutil.copy2(src, out)

wb = openpyxl.load_workbook(out)

# Number formats showing "-" for zero:
FMT_INT_DASH = '#,##0;-#,##0;"-"'
FMT_DEC2_DASH = '#,##0.00;-#,##0.00;"-"'
FMT_PCT_DASH = '0.0%;-0.0%;"-"'

for ws in wb.worksheets:
    name = ws.title
    max_row = ws.max_row
    max_col = ws.max_column
    
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            
            # Skip non-numeric, empty, or string cells
            if val is None or isinstance(val, str):
                continue
            if not isinstance(val, (int, float)):
                continue
            
            # Sheet 5: 비중 columns (every other column starting from col 4 for seasons)
            if name == '5.시즌공장' and c >= 4:
                rel = (c - 4) % 2
                if rel == 1:  # 비중 column (odd offset from col 4)
                    # Convert ratio to percentage value and format
                    ws.cell(row=r, column=c).value = round(val * 100, 1)
                    cell.number_format = '0.0'
                    if val == 0:
                        cell.value = '-'
                    continue
            
            # Sheet 4: 비중% columns (every other column for seasons: 26 SPR=col 4-5, 26 SUM=col 6-7, etc.)
            if name == '4.시즌부서공장' and c >= 5:
                rel = (c - 5) % 2
                if rel == 0:  # 비중% column (cols 5, 7, 9, 11)
                    cell.number_format = '0.0'
                    if val == 0:
                        cell.value = '-'
                    continue
            
            # All other number cells: apply comma format with dash for zero
            if isinstance(val, float) and val != int(val):
                cell.number_format = FMT_DEC2_DASH
            else:
                cell.number_format = FMT_INT_DASH

wb.save(out)
import os; print(f'Saved: {out} ({os.path.getsize(out)/1024:.0f} KB)')
wb.close()
