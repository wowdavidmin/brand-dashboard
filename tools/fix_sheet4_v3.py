#!/usr/bin/env python3
"""Rebuild Sheet 4: dept | brand | factory | season | qty | share"""
import openpyxl, os, shutil
from openpyxl.styles import Font, Border, Side, Alignment
from collections import defaultdict

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
src = os.path.join(BASE, '리포트_5Sheet_최종_v3_20260630.xlsx')
out = os.path.join(BASE, '리포트_5Sheet_최종_v6_20260630.xlsx')
shutil.copy2(src, out)

wb = openpyxl.load_workbook(out)
ws4 = wb['4.시즌부서공장']
rows = list(ws4.iter_rows(values_only=True))

# Brand names from row 2, dept names from row 1
brand_to_dept = {}
last_dept = ''
for c in range(2, len(rows[2])):
    bn = str(rows[2][c] or '').strip()
    dn = str(rows[1][c] or '').strip() if c < len(rows[1]) else ''
    if dn: last_dept = dn
    if bn and bn not in ('', 'Factory'):
        brand_to_dept[bn] = last_dept

# Parse all (dept, brand, factory, season) -> qty
data = defaultdict(int)
cur_season = ''
for ri in range(3, len(rows)):
    row = rows[ri]
    s = str(row[0] or '').strip()
    if s in ('26 SPR','26 SUM','26 FAL','26 HOL'): cur_season = s
    fac = str(row[1] or '').strip()
    if not fac or not cur_season: continue
    for c in range(2, len(row)):
        bn = str(rows[2][c] or '').strip()
        if not bn or bn in ('', 'Factory'): continue
        q = row[c] if row[c] else 0
        try: q = int(float(str(q)))
        except: q = 0
        if q > 0:
            dn = brand_to_dept.get(bn, 'BU4부')
            data[(dn, bn, fac, cur_season)] += q

del wb['4.시즌부서공장']
ws4n = wb.create_sheet('4.시즌부서공장')

thin = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
hdr = Font(bold=True, size=10, name='Calibri')

ws4n.cell(row=1, column=1, value='2026 Season 조직-브랜드-Factory-시즌별 배정 수량 및 비중').font = Font(bold=True, size=13, name='Calibri')
ws4n.merge_cells('A1:F1')

cols = ['조직', '브랜드', 'Factory', '시즌', '수량', '비중']
for ci, h in enumerate(cols, 1):
    ws4n.cell(row=3, column=ci, value=h).font = hdr

FMT = '#,##0'
r = 4
# Group by (dept, brand, factory) for share calculation
grouped = defaultdict(lambda: defaultdict(int))
for (dn, bn, fac, s), q in data.items():
    grouped[(dn, bn, fac)][s] += q

for (dn, bn, fac), seasons_qty in sorted(grouped.items()):
    total_qty = sum(seasons_qty.values())
    for s in ('26 SPR', '26 SUM', '26 FAL', '26 HOL'):
        qty = seasons_qty.get(s, 0)
        share = round(qty / total_qty, 4) if total_qty else 0
        
        c = 1
        for v in [dn, bn, fac, s]:
            cell = ws4n.cell(row=r, column=c, value=v)
            cell.font = Font(size=10); cell.border = thin; cell.alignment = Alignment(horizontal='left'); c += 1
        
        cell_q = ws4n.cell(row=r, column=c, value=qty if qty else 0)
        cell_q.number_format = FMT; cell_q.font = Font(size=10); cell_q.border = thin
        cell_q.alignment = Alignment(horizontal='right'); c += 1
        
        cell_p = ws4n.cell(row=r, column=c, value=share if qty else 0)
        cell_p.number_format = '0.0000'; cell_p.font = Font(size=10); cell_p.border = thin
        cell_p.alignment = Alignment(horizontal='right')
        
        r += 1

# Fix column widths
ws4n.column_dimensions['A'].width = 16
ws4n.column_dimensions['B'].width = 22
ws4n.column_dimensions['C'].width = 32
ws4n.column_dimensions['D'].width = 10
ws4n.column_dimensions['E'].width = 12
ws4n.column_dimensions['F'].width = 10

order = ['1.시즌현황','2.시즌별','3.월별','4.시즌부서공장','5.시즌공장']
for i, name in enumerate(order):
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        if idx != i: wb.move_sheet(name, offset=i-idx)

wb.save(out)
import os; print(f'Saved: {out} ({os.path.getsize(out)/1024:.0f} KB)')
wb.close()
