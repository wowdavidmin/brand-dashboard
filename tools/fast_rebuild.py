#!/usr/bin/env python3
"""Fast rebuild: use data from old Sheet 1, rearrange into sections"""
import openpyxl, os, shutil
from openpyxl.styles import Font, Border, Side, Alignment

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
TEMP = os.path.join(BASE, '_temp_sheets2to5.xlsx')
OUT = os.path.join(BASE, '리포트_5Sheet_최종_20260630.xlsx')

# Copy temp to output
shutil.copy2(TEMP, OUT)
print('Copied temp to output')

# Read old Sheet 1 data
wb = openpyxl.load_workbook(OUT)
ws_old = wb['1.시즌현황']

# Read brand names and data
brand_cols = {}
all_data = {}  # (season, brand) -> [qty, amt, net, up, margin, cm]
seasons_list = ['26 SPR', '26 SUM', '26 FAL', '26 HOL']

for row in ws_old.iter_rows(min_row=1, max_row=5, values_only=False):
    r = row[0].row
    vals = [c.value for c in row]
    if r == 4:  # Brand names
        for c, v in enumerate(vals):
            if c >= 1 and v and isinstance(v, str) and v.strip():
                brand_cols[c] = v.strip()
    elif r == 5:  # Metric labels
        pass

# Read data rows (r=6 to 9 = seasons, r=10 = Total)
for r in range(6, 11):
    row_vals = [c.value for c in ws_old[r]]
    season = str(row_vals[0] or '').strip() if row_vals else ''
    if not season: continue
    
    for c, brand in brand_cols.items():
        # c to c+5 are: 수량, 총매출, 순매출, 평균단가, 마진율, CM
        qty = row_vals[c] if c < len(row_vals) else 0
        amt = row_vals[c+1] if c+1 < len(row_vals) else 0
        net = row_vals[c+2] if c+2 < len(row_vals) else 0
        up = row_vals[c+3] if c+3 < len(row_vals) else 0
        margin = row_vals[c+4] if c+4 < len(row_vals) else 0
        cm = row_vals[c+5] if c+5 < len(row_vals) else 0
        all_data[(season, brand)] = [qty, amt, net, up, margin, cm]

print(f'Read data for {len(all_data)} entries, {len(brand_cols)} brands')

# Delete old Sheet 1 and create new one
del wb['1.시즌현황']
ws1 = wb.create_sheet('1.시즌현황', 0)

# Title
ws1.cell(row=1, column=1, value='[약진통상] 2026 Season 부서별 브랜드별 현황 (단위: $K)').font = Font(bold=True, size=13, name='Calibri')
ws1.merge_cells('A1:ZZ1')

# Brand department ordering
REF = [
    ('BU1-영업1부', ['GAP', 'IPSS GAP', 'BRFS', 'Madhappy', 'Mott & Bow']),
    ('BU2-영업1부', ['OLD NAVY', "ON  MEN'S"]),
    ('BU2-영업2부', ['WALMART', 'QUINCE']),
    ('BU3-영업1부', ['BR', 'IPSS ON']),
    ('BU3-영업2부', ["KOHL'S", 'URBAN']),
    ('BU3-영업3부', ['AERIE', 'Offline']),
    ('BU3-영업4부', ['MLB']),
    ('BU4부', ['NUUDS', 'ONER ACTIVE', 'ROSS', 'AEO']),
    ('BU5-영업5부', ['Justice', 'INGRID & ISABEL']),
    ('GTBU-과테말라', ['OLDNAVY(HYBRID)']),
]

# Which brands are active (have data)
active = [(dn, b) for dn, brs in REF for b in brs 
          if any(all_data.get((s, b), [0]*6)[0] for s in seasons_list)]

thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
hdr = Font(bold=True, size=10, name='Calibri')

metrics = [
    ('1. ORDER 수량', 0, '#,##0'),
    ('2. 총매출액', 1, '#,##0'),
    ('3. 순매출액', 2, '#,##0'),
    ('4. 평균단가', 3, '#,##0.00'),
    ('5. 마진율', 4, '0.00'),
    ('6. CM', 5, '#,##0.00'),
]

def safe_val(v):
    if v is None: return 0
    if isinstance(v, (int, float)): return v
    return 0

r = 3
for label, midx, fmt in metrics:
    # Section title
    ws1.cell(row=r, column=1, value=label).font = Font(bold=True, size=12, name='Calibri')
    r += 1
    
    # Dept headers
    col = 2
    for dn, brs in REF:
        these = [(d,b) for (d,b) in active if d == dn]
        if not these: continue
        cs = col; col += len(these)
        ws1.merge_cells(start_row=r, start_column=cs, end_row=r, end_column=col-1)
        ws1.cell(row=r, column=cs, value=dn).font = hdr
        ws1.cell(row=r, column=cs).alignment = Alignment(horizontal='center')
    r += 1
    
    # Brand headers
    ws1.cell(row=r, column=1, value='시즌').font = hdr
    ws1.cell(row=r, column=1).border = thin
    col = 2
    for dn, b in active:
        cell = ws1.cell(row=r, column=col, value=b)
        cell.font = Font(bold=True, size=8, name='Calibri')
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
        cell.border = thin
        col += 1
    r += 1
    
    # Data rows
    for season in seasons_list + ['Total']:
        ws1.cell(row=r, column=1, value=season).font = Font(bold=(season=='Total'), size=10, name='Calibri')
        ws1.cell(row=r, column=1).border = thin
        col = 2
        for dn, b in active:
            vals = all_data.get((season, b), [0]*6)
            v = safe_val(vals[midx]) if midx < len(vals) else 0
            cell = ws1.cell(row=r, column=col, value=v)
            cell.number_format = fmt if midx != 4 else '0.00'
            cell.font = Font(bold=(season=='Total'), size=10, name='Calibri')
            cell.border = thin
            cell.alignment = Alignment(horizontal='right')
            col += 1
        r += 1
    r += 1  # blank separator

wb.save(OUT)
print(f'Saved: {OUT}')
import os; print(f'Size: {os.path.getsize(OUT)/1024:.0f} KB')
print('Sheets:', wb.sheetnames)
wb.close()
