#!/usr/bin/env python3
"""Rebuild Sheet 3: brands as rows, months as columns"""
import openpyxl, os
from openpyxl.styles import Font, Border, Side, Alignment

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
f = os.path.join(BASE, '리포트_5Sheet_최종_v2_20260630.xlsx')

wb = openpyxl.load_workbook(f)
ws3 = wb['3.월별']
rows = list(ws3.iter_rows(values_only=True))

print(f'Total rows read: {len(rows)}')
for ri, row in enumerate(rows):
    v0 = str(row[0])[:30] if row[0] is not None else '(empty)'
    v1 = str(row[1])[:15] if len(row) > 1 and row[1] is not None else ''
    v4 = str(row[4])[:15] if len(row) > 4 and row[4] is not None else ''
    print(f'  [{ri}] {v0:30s} | {v1:15s} | {v4:15s}')

# Find brand name row and data rows
brand_row_idx = None
data_start = None
for ri, row in enumerate(rows):
    v0 = str(row[0] or '').strip()
    if v0 == '월':
        brand_row_idx = ri
    elif v0.endswith('월'):
        if data_start is None:
            data_start = ri

print(f'\nBrand row index: {brand_row_idx}, Data start: {data_start}')

if brand_row_idx is None:
    print('ERROR: Could not find brand row!')
    exit(1)

# Parse brand names from brand row (cols 1, 4, 7, ...)
brand_row = rows[brand_row_idx]
brand_names = []
for c in range(1, len(brand_row), 3):
    bn = str(brand_row[c] or '').strip()
    if bn and bn not in ('수량', '매출', '순매출'):
        brand_names.append(bn)

print(f'Brands: {brand_names}')

# Parse month data
data = {}
months_list = []
for ri in range(data_start, len(rows)):
    row = rows[ri]
    mn = str(row[0] or '').strip()
    if not mn or not mn.endswith('월'):
        continue
    months_list.append(mn)
    for bi, bn in enumerate(brand_names):
        c = 1 + bi * 3
        def sv(v):
            if v is None: return 0
            if isinstance(v, (int, float)): return int(v)
            try: return int(float(str(v)))
            except: return 0
        q = sv(row[c]) if c < len(row) else 0
        s = sv(row[c+1]) if c+1 < len(row) else 0
        n = sv(row[c+2]) if c+2 < len(row) else 0
        data[(mn, bn)] = [q, s, n]

print(f'Months: {months_list}')
print(f'Data entries: {len(data)}')
print(f'Sample: GAP 1월 = {data.get(("1월", "GAP"))}')

# Delete old, create new
del wb['3.월별']
ws3n = wb.create_sheet('3.월별')

thin = Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
hdr = Font(bold=True, size=10, name='Calibri')

ws3n.cell(row=1, column=1, value='2026년 월별 부서별 브랜드별 매출현황 (단위: $K)').font = Font(bold=True, size=13, name='Calibri')
ws3n.merge_cells('A1:ZZ1')

col = 2
mc = {}
for mn in months_list:
    ws3n.cell(row=3, column=col, value=mn).font = hdr
    ws3n.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+2)
    ws3n.cell(row=4, column=col, value='수량').font = hdr
    ws3n.cell(row=4, column=col+1, value='매출').font = hdr
    ws3n.cell(row=4, column=col+2, value='순매출').font = hdr
    mc[mn] = col; col += 3

ws3n.cell(row=3, column=1, value='브랜드/부서').font = hdr

REF = [
    ('BU1-영업1부', ['GAP', 'IPSS GAP', 'BRFS', 'Madhappy', 'Mott & Bow']),
    ('BU2-영업1부', ['OLD NAVY', "ON  MEN'S"]),
    ('BU2-영업2부', ['WALMART', 'QUINCE']),
    ('BU3-영업1부', ['BR', 'IPSS ON']),
    ('BU3-영업2부', ["KOHL'S", 'URBAN']),
    ('BU3-영업3부', ['AERIE', 'Offline']),
    ('BU3-영업4부', ['MLB']),
    ('BU4부', ['NUUDS', 'ONER ACTIVE', 'ROSS', 'AEO']),
    ('BU5-영업5부', ['Justice', 'INGRID & ISABEL']),
    ('GTBU-과테말라', ['OLDNAVY(HYBRID)']),
]

FMT = '#,##0'
r = 5
for dn, brs in REF:
    active = [b for b in brs if any(data.get((mn, b), [0]*3)[0] for mn in months_list)]
    if not active: continue
    for b in active:
        ws3n.cell(row=r, column=1, value=b)
        for mn in months_list:
            c = mc[mn]; d = data.get((mn, b), [0, 0, 0])
            for ci, v in enumerate([d[0], d[1], d[2]]):
                cell = ws3n.cell(row=r, column=c+ci, value=v)
                cell.number_format = FMT; cell.font = Font(size=10); cell.border = thin
                cell.alignment = Alignment(horizontal='right')
        ws3n.cell(row=r, column=1).alignment = Alignment(horizontal='left')
        r += 1
    ws3n.cell(row=r, column=1, value=f'{dn} 요약').font = Font(bold=True)
    for mn in months_list:
        c = mc[mn]; dq = sum(data.get((mn, b), [0]*3)[0] for b in active)
        da = sum(data.get((mn, b), [0]*3)[1] for b in active)
        dn_ = sum(data.get((mn, b), [0]*3)[2] for b in active)
        for ci, v in enumerate([dq, da, dn_]):
            cell = ws3n.cell(row=r, column=c+ci, value=v)
            cell.font = Font(bold=True); cell.border = thin; cell.alignment = Alignment(horizontal='right')
    r += 1

# Fix sheet order
order = ['1.시즌현황', '2.시즌별', '3.월별', '4.시즌부서공장', '5.시즌공장']
for i, name in enumerate(order):
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        if idx != i:
            wb.move_sheet(name, offset=i-idx)

out = os.path.join(BASE, '리포트_5Sheet_최종_v3_20260630.xlsx')
wb.save(out)
print(f'\nSaved: {out} ({os.path.getsize(out)/1024:.0f} KB)')
wb.close()
