#!/usr/bin/env python3
import openpyxl

BASE = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week'
f = BASE + r'\영업DATA_5Sheet_FINAL_20260630.xlsx'
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb['1.시즌현황']

brands = {}
total_row = None
for row in ws.iter_rows(values_only=True):
    if row[0] == '시즌':
        for c, v in enumerate(row):
            if c >= 1 and v and isinstance(v, str) and v not in ('수량','총매출','순매출','평균단가','마진율%','CM%',''):
                brands[c] = v.strip()
    if row[0] == 'Total':
        total_row = row
        break

print('=== BU1-영업1부 Total ===')
bu1_brands = ['GAP', 'IPSS GAP', 'BRFS', 'Madhappy', 'Mott&Bow']
bu1_total = 0
for c in sorted(brands.keys()):
    b = brands[c]
    v = total_row[c] if total_row and total_row[c] else 0
    if not isinstance(v, (int, float)): v = 0
    print(f'  {b:20s}: {int(v):>10,}')
    if b in bu1_brands:
        bu1_total += int(v)

print(f'  {"-"*32}')
print(f'  {"BU1 합계":20s}: {bu1_total:>10,}')
print(f'  {"참고(2주전)":20s}: 28,287,329')
print(f'  {"증감":20s}: {bu1_total-28287329:>+10,}')
if 28287329:
    print(f'  {"증감률":20s}: {(bu1_total-28287329)/28287329*100:+.1f}%')

# Also check BU2
print()
print('=== BU2-영업1부 (OLD NAVY + ON) ===')
bu2_brands = ['OLD NAVY', "ON MEN'S"]
bu2_total = 0
for c in sorted(brands.keys()):
    b = brands[c]
    v = total_row[c] if total_row and total_row[c] else 0
    if not isinstance(v, (int, float)): v = 0
    if b in bu2_brands:
        bu2_total += int(v)
        print(f'  {b:20s}: {int(v):>10,}')
print(f'  {"BU2-1 합계":20s}: {bu2_total:>10,}')
print(f'  {"참고":20s}: 64,810,266 (ON MEN/WMN/KID 포함)')

print()
print('=== 시트 구성 ===')
for s in wb.sheetnames:
    ws2 = wb[s]
    cnt = sum(1 for _ in ws2.iter_rows(values_only=True))
    print(f'  [{s}] {cnt}행')

wb.close()
