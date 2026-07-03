import openpyxl

# v12 original file
v12 = r'C:\Users\user\.openclaw\media\inbound\영업데이터파일_v12_20260702---c5b345bd-47c5-47cf-b54c-e337523c2574.xlsx'
wb = openpyxl.load_workbook(v12)
ws = wb.worksheets[0]

# Pivot analysis file
pivot = r'C:\Users\user\.openclaw\workspace\브랜드별_직접원가_피벗분석_v1.0_20260629.xlsx'
pwb = openpyxl.load_workbook(pivot)
pws1 = pwb['1.시즌별 부서브랜드 원가']

# === Step 1: Build pivot lookup: (brand, season) -> PO금액(C4), PO수량(C7), Profit(C6), 직접원가(C5) ===
pivot_data = {}  # key: (brand, season) -> {'PO금액': , 'PO수량': , 'PROFIT': , '직접원가': }
for r in range(2, pws1.max_row+1):
    dept = str(pws1.cell(r, 1).value or '').strip()
    brand = str(pws1.cell(r, 2).value or '').strip()
    season = str(pws1.cell(r, 3).value or '').strip()
    
    if brand in ('소계', '총계') or not season.startswith('26 '):
        continue
    
    po_amt = pws1.cell(r, 4).value or 0  # PO금액
    direct_cost = pws1.cell(r, 5).value or 0
    profit = pws1.cell(r, 6).value or 0
    po_qty = pws1.cell(r, 7).value or 0  # PO수량
    
    key = (brand, season)
    if key not in pivot_data:
        pivot_data[key] = {'PO금액': 0, 'PO수량': 0, 'PROFIT': 0, '직접원가': 0}
    
    pivot_data[key]['PO금액'] += float(po_amt if po_amt else 0)
    pivot_data[key]['PO수량'] += float(po_qty if po_qty else 0)
    pivot_data[key]['PROFIT'] += float(profit if profit else 0)
    pivot_data[key]['직접원가'] += float(direct_cost if direct_cost else 0)

# === Step 2: Compare with v12 file for BU3 brands ===
# BU3-영업1부: BR (C15), IPSS ON (C16)
# BU3-영업2부: KOHL'S (C18), URBAN (C19)
# BU3-영업3부: AERIE (C21), Offline (C22)
# BU3-영업4부: MLB (C24)

# Season mapping
sections = {
    'ORDER': {'header_row': 3, 'data': {6: '26 SPR', 7: '26 SUM', 8: '26 FAL', 9: '26 HOL', 10: 'Total'}, 'label': 'ORDER 수량'},
    '총매출액': {'header_row': 12, 'data': {15: '26 SPR', 16: '26 SUM', 17: '26 FAL', 18: '26 HOL', 19: 'Total'}, 'label': '총매출액'},
    '순매출액': {'header_row': 21, 'data': {24: '26 SPR', 25: '26 SUM', 26: '26 FAL', 27: '26 HOL', 28: 'Total'}, 'label': '순매출액'},
}

bu3_brands = {
    15: 'BANANA REPUBLIC',  # BR
    16: 'ON IPSS',          # IPSS ON
    18: "KOHL'S",           # KOHL'S
    19: 'URBAN',            # URBAN
    21: 'AERIE',            # AERIE
    22: 'Offline',          # Offline
    24: 'MLB',              # MLB
}

print('=== BU3 브랜드: v12 vs pivot (PO금액) 비교 ===')
print()

# Compare 총매출액 (which should = PO금액 in $K)
for col, brand in bu3_brands.items():
    print(f'--- {brand} (C{col}) ---')
    for row, season in sections['총매출액']['data'].items():
        v12_val = ws.cell(row, col).value
        # Try '-' to 0
        if v12_val == '-' or v12_val is None:
            v12_val = 0
        
        # Look up pivot
        pivot_key = (brand, season)
        if pivot_key in pivot_data:
            pivot_po = pivot_data[pivot_key]['PO금액']  # in dollars
            pivot_po_k = round(pivot_po / 1000, 1)  # convert to $K
            v12_float = float(v12_val) if v12_val != '-' and v12_val is not None else 0
            diff = v12_float - pivot_po_k
            print(f'  {season}: v12={v12_float:>10.1f} $K | pivot={pivot_po_k:>10.1f} $K | diff={diff:>+10.1f} $K')
        else:
            print(f'  {season}: v12={v12_val} | pivot=N/A')
    print()

# Also compare ORDER 수량
print()
print('=== BU3 브랜드: v12 vs pivot (PO수량) 비교 ===')
print()
for col, brand in bu3_brands.items():
    print(f'--- {brand} (C{col}) ---')
    for row, season in sections['ORDER']['data'].items():
        v12_val = ws.cell(row, col).value
        if v12_val == '-' or v12_val is None:
            v12_val = 0
        
        pivot_key = (brand, season)
        if pivot_key in pivot_data:
            pivot_qty = pivot_data[pivot_key]['PO수량']
            v12_float = float(v12_val) if v12_val != '-' and v12_val is not None else 0
            diff = v12_float - pivot_qty
            print(f'  {season}: v12={v12_float:>12.0f} | pivot={pivot_qty:>12.0f} | diff={diff:>+12.0f}')
        else:
            print(f'  {season}: v12={v12_val} | pivot=N/A')
    print()

pwb.close()
wb.close()
