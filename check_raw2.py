import openpyxl

# June 9 raw data files
f1 = r'C:\Users\user\Downloads\raw data 검증_20260609.xlsx'
f2 = r'C:\Users\user\Downloads\raw_data_중복제거_20260609.xlsx'
f3 = r'C:\Users\user\Downloads\브랜드별 직접원가_20260609.xlsx'

for fpath in [f1, f2]:
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        print(f'=== {fpath.split(chr(92))[-1]} ===')
        print(f'  Sheets: {wb.sheetnames[:5]}...')
        ws = wb.worksheets[0]
        print(f'  Rows: {ws.max_row}, Cols: {ws.max_column}')
        # Print first 5 rows first 15 cols
        for r in range(1, min(6, ws.max_row+1)):
            row_vals = []
            for c in range(1, min(16, ws.max_column+1)):
                v = ws.cell(r, c).value
                if v is not None:
                    row_vals.append(f'C{c}={v}')
            if row_vals:
                print(f'  Row {r}: {", ".join(str(x) for x in row_vals)}')
        print()
        wb.close()
    except Exception as e:
        print(f'Error with {fpath}: {e}')
        print()

# Also check the 브랜드별 직접원가 file
try:
    wb = openpyxl.load_workbook(f3, read_only=True, data_only=True)
    print(f'=== {f3.split(chr(92))[-1]} ===')
    print(f'  Sheets: {wb.sheetnames[:5]}...')
    ws = wb.worksheets[0]
    print(f'  Rows: {ws.max_row}, Cols: {ws.max_column}')
    for r in range(1, min(6, ws.max_row+1)):
        row_vals = []
        for c in range(1, min(20, ws.max_column+1)):
            v = ws.cell(r, c).value
            if v is not None:
                row_vals.append(f'C{c}={v}')
        if row_vals:
            print(f'  Row {r}: {", ".join(str(x) for x in row_vals)}')
    wb.close()
except Exception as e:
    print(f'Error with {f3}: {e}')
