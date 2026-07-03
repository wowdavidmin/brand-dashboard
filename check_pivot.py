import openpyxl

# The existing pivot analysis file
pivot = r'C:\Users\user\.openclaw\workspace\브랜드별_직접원가_피벗분석_v1.0_20260629.xlsx'

wb = openpyxl.load_workbook(pivot)
print(f'Sheets: {wb.sheetnames}')
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'\n=== Sheet: {sn} (rows={ws.max_row}, cols={ws.max_column}) ===')
    for r in range(1, min(ws.max_row+1, 20)):
        row_vals = []
        for c in range(1, min(ws.max_column+1, 20)):
            v = ws.cell(r, c).value
            if v is not None:
                row_vals.append(f'C{c}={v}')
        if row_vals:
            print(f'  Row {r}: {", ".join(str(x) for x in row_vals)}')
wb.close()
