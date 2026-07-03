import openpyxl

src = r'C:\Users\user\.openclaw\workspace\영업데이터파일_v12_20260702_processed.xlsx'
wb = openpyxl.load_workbook(src)
ws = wb.worksheets[0]

# 마진율 section: rows 42-46, columns 3-35
for r in range(42, 47):
    for c in range(3, 36):
        cell = ws.cell(r, c)
        if cell.value is not None and cell.value != '-' and cell.value != '':
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.0"%"'

out1 = src
out2 = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week\영업데이터파일_v12_20260702.xlsx'
out3 = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week\영업데이터파일_v13_20260702.xlsx'

wb.save(out1)
wb.save(out2)
wb.save(out3)

# Verify
wb2 = openpyxl.load_workbook(out1)
ws2 = wb2.worksheets[0]
print('마진율 확인 (26 SPR):')
print(f'  GAP: {ws2.cell(42, 3).value} (format: {ws2.cell(42, 3).number_format})')
print(f'  합계: {ws2.cell(42, 35).value} (format: {ws2.cell(42, 35).number_format})')
print(f'Total GAP: {ws2.cell(46, 3).value} (format: {ws2.cell(46, 3).number_format})')
