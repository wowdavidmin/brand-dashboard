import openpyxl

pivot = r'C:\Users\user\.openclaw\workspace\브랜드별_직접원가_피벗분석_v1.0_20260629.xlsx'
pwb = openpyxl.load_workbook(pivot)
pws = pwb['1.시즌별 부서브랜드 원가']

print('=== All departments in pivot ===')
depts = set()
for r in range(2, pws.max_row+1):
    dept = str(pws.cell(r, 1).value or '').strip()
    brand = str(pws.cell(r, 2).value or '').strip()
    season = str(pws.cell(r, 3).value or '').strip()
    if brand not in ('소계', '총계'):
        depts.add(dept)

for d in sorted(depts):
    print(f'  {d}')

print()

# Print all rows for BU3-related departments
print('=== 영업3부 or BU3 related rows ===')
for r in range(2, pws.max_row+1):
    dept = str(pws.cell(r, 1).value or '').strip()
    brand = str(pws.cell(r, 2).value or '').strip()
    season = str(pws.cell(r, 3).value or '').strip()
    
    if (('3부' in dept or '3_D' in dept) or 
        brand in ('BANANA REPUBLIC', 'ON IPSS', "KOHL`S", 'URBAN OUTFITTERS', 'AERIE', 'Offline', 'MLB')):
        
        amt = pws.cell(r, 4).value or 0
        qty = pws.cell(r, 7).value or 0
        print(f'  {dept:15s} | {brand:25s} | {season:10s} | amt={float(amt)/1000:>8.1f}K | qty={float(qty):>8.0f}')

pwb.close()
