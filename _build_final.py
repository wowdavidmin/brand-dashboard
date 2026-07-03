"""
Raw data -> 4 pivot sheets matching template format (1-4)
"""
import pandas as pd
import numpy as np
import warnings, os
warnings.filterwarnings('ignore')

SRC = r'C:\Users\user\.openclaw\media\inbound\브랜드별_직접원가_raw_v1.0_20260629---b99ec911-9409-4c91-ae86-0a63ab160c8c.xlsx'
OUT = r'C:\Users\user\.openclaw\workspace\브랜드별_직접원가_피벗분석_v2.0_20260629.xlsx'

print("Reading raw data...")
df = pd.read_excel(SRC, sheet_name='Raw_20260629')
print(f"Shape: {df.shape}")

# Clean columns
df.columns = [c.replace('\n', ' ').strip() for c in df.columns]

# Convert numeric
num_cols = ['PO수량', 'PO금액', '매출할인', '원자재비 원사', '원자재비 원단', '부자재비 본사',
            '염가공', 'I.F염가공', '염가공비 계', '임가공료 외화', '임가공료 원화', '임가공료 합계',
            'CMT 봉제료', 'CMT I.F', 'CMT 부자재', 'CMT 임가공료해외', 'CMT 기타수수료', 'CMT 계',
            '해상운임', 'Inspection Fee', 'NCPM', '직접원가 AMOUNT', '직접원가 %',
            'PROFIT AMOUNT', 'PROFIT %']
for c in num_cols:
    if c in df.columns:
        df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0)

# Clean text cols - handle mixed types safely
def safe_str(x):
    if pd.isna(x):
        return ''
    return str(x).strip()

for c in ['브랜드', '부서', '시즌', 'Contract Factory']:
    if c in df.columns:
        df[c] = df[c].apply(safe_str)

# Remove empty brands/seasons
df = df[df['브랜드'] != ''].copy()
df = df[df['시즌'] != ''].copy()

brands = sorted(df['브랜드'].unique().tolist())
depts = sorted(df['부서'].unique().tolist())
seasons = sorted(df['시즌'].unique().tolist())
print(f"Brands ({len(brands)}): {brands}")
print(f"Depts ({len(depts)}): {depts}")
print(f"Seasons ({len(seasons)}): {seasons}")

# ============================================
# SHEET 1: 시즌별 부서/브랜드 직접원가 비교
# Like template Sheet 1: hierarchical dept > brand with metrics
# ============================================
print("\nBuilding Sheet 1...")
dept_label_map = {
    'BU1_D1': '영업1부', 'BU1_D2': '영업1부',
    'BU2_D1': '영업2부', 'BU2_D2': '영업2부',
    'BU3_D1': '영업3부', 'BU3_D2': '영업3부',
    'BU3_D3': '영업3부', 'BU3_D4': '영업3부',
    'BU4': '영업4부', 'BU5': '영업5부',
    'GTBU': 'GTBU',
    '인터텍스 기획부': '인터텍스'
}
df['부서명'] = df['부서'].map(dept_label_map).fillna(df['부서'])

s1 = df.groupby(['부서명', '브랜드', '시즌']).agg({
    'PO금액': 'sum',
    '직접원가 AMOUNT': 'sum',
    'PROFIT AMOUNT': 'sum',
    'PO수량': 'sum',
}).reset_index()
mask = s1['PO금액'] > 0
s1['직접원가율(%)'] = np.where(mask, (s1['직접원가 AMOUNT'] / s1['PO금액'] * 100).round(2), 0)
s1['마진율(%)'] = np.where(mask, (s1['PROFIT AMOUNT'] / s1['PO금액'] * 100).round(2), 0)

# ============================================
# SHEET 2: 브랜드별 원가구성 분석
# Like template Sheet 2: brand별 cost breakdown
# ============================================
print("Building Sheet 2...")
s2 = df.groupby(['브랜드']).agg({
    'PO금액': 'sum',
    '원자재비 원사': 'sum',
    '원자재비 원단': 'sum',
    '부자재비 본사': 'sum',
    '염가공비 계': 'sum',
    '임가공료 합계': 'sum',
    'CMT 계': 'sum',
    '해상운임': 'sum',
    '직접원가 AMOUNT': 'sum',
    'PROFIT AMOUNT': 'sum',
    'PO수량': 'sum',
}).reset_index()

m2 = s2['PO금액'] > 0
s2['직접원가율(%)'] = np.where(m2, (s2['직접원가 AMOUNT'] / s2['PO금액'] * 100).round(2), 0)
s2['마진율(%)'] = np.where(m2, (s2['PROFIT AMOUNT'] / s2['PO금액'] * 100).round(2), 0)

cost_cols = ['원자재비 원사', '원자재비 원단', '부자재비 본사', '염가공비 계', '임가공료 합계', 'CMT 계', '해상운임']
for col in cost_cols:
    s2[f'{col} 구성비(%)'] = np.where(s2['직접원가 AMOUNT'] > 0,
        (s2[col] / s2['직접원가 AMOUNT'] * 100).round(2), 0)

# ============================================
# SHEET 3: 시즌별 브랜드 Pivot  
# Like template Sheet 3: season x brand pivot with PO금액
# ============================================
print("Building Sheet 3...")
season_order = ['24 SPR','24 SUM','24 FAL','24 HOL',
                '25 SPR','25 SUM','25 FAL','25 HOL',
                '26 BTS','26 SPR','26 SUM','26 FAL','26 HOL','27 SPR']
existing = [s for s in season_order if s in df['시즌'].unique()]
others = [s for s in df['시즌'].unique() if s not in season_order and s != '']

# Pivot: 시즌 x 브랜드 values=PO금액
p3 = df.pivot_table(index='시즌', columns='브랜드', values='PO금액', aggfunc='sum', fill_value=0)
p3 = p3.reindex(existing + others)
p3 = p3.fillna(0)
p3['TOTAL'] = p3.sum(axis=1)
p3 = p3.reset_index()

# ============================================
# SHEET 4: 공장(Contract Factory)별 원가 분석
# ============================================
print("Building Sheet 4...")
s4 = df.groupby(['Contract Factory']).agg({
    'PO금액': 'sum',
    'CMT 계': 'sum',
    '해상운임': 'sum',
    '직접원가 AMOUNT': 'sum',
    'PROFIT AMOUNT': 'sum',
    'PO수량': 'sum',
}).reset_index()

m4 = s4['PO금액'] > 0
s4['직접원가율(%)'] = np.where(m4, (s4['직접원가 AMOUNT'] / s4['PO금액'] * 100).round(2), 0)
s4['마진율(%)'] = np.where(m4, (s4['PROFIT AMOUNT'] / s4['PO금액'] * 100).round(2), 0)
s4['CMT비중(%)'] = np.where(s4['직접원가 AMOUNT'] > 0,
    (s4['CMT 계'] / s4['직접원가 AMOUNT'] * 100).round(2), 0)
s4 = s4.sort_values('PO금액', ascending=False)

# ============================================
# Write to Excel
# ============================================
print("Writing Excel...")
with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
    s1.to_excel(writer, sheet_name='1.시즌별 부서브랜드 원가', index=False)
    s2.to_excel(writer, sheet_name='2.브랜드별 원가구성 분석', index=False)
    p3.to_excel(writer, sheet_name='3.시즌별 브랜드 Pivot', index=False)
    s4.to_excel(writer, sheet_name='4.공장별 원가 분석', index=False)

# Formatting
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = load_workbook(OUT)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=10)

for ws in wb.worksheets:
    for cell in ws[1]:
        if cell.value:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

wb.save(OUT)
print(f"Done! Saved to: {OUT}")
print(f"  Sheet 1: {len(s1)} rows")
print(f"  Sheet 2: {len(s2)} rows")
print(f"  Sheet 3: {len(p3)} seasons x {len(p3.columns)} cols")
print(f"  Sheet 4: {len(s4)} factories")
