import openpyxl

src = r'C:\Users\user\.openclaw\workspace\영업데이터파일_v12_20260702_processed.xlsx'
wb = openpyxl.load_workbook(src)
ws = wb.worksheets[0]

# Quick verify key BU3 values
print('KOHL_S (C18):')
for r in [6,7,8,9,10]:
    s = ws.cell(r,2).value
    print(f'  {s}: ORDER={ws.cell(r,18).value}')

print()
print('URBAN (C19):')
for r in [6,7,8,9,10]:
    s = ws.cell(r,2).value  
    print(f'  {s}: ORDER={ws.cell(r,19).value}')

print()
print('S.TOTAL BU3-2bu (C20):')
for r in [6,7,8,9,10]:
    s = ws.cell(r,2).value
    print(f'  {s}: ORDER={ws.cell(r,20).value}')

# Compare with expected values from pivot
print()
print('Expected (from pivot):')
print('KOHL_S: SPR=2573501, SUM=3321230, FAL=5944313, HOL=2903050')
print('URBAN: SPR=6400, SUM=11625, FAL=21220, HOL=0')
print('S.TOTAL: SPR=2579901, SUM=3332855, FAL=5965533, HOL=2903050')
