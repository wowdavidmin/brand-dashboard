import openpyxl

# File 1: RATIO 등록 현황
path1 = r"C:\Users\user\.openclaw\media\inbound\RATIO_등록_현황_1월_6월---cbfe54da-7d01-4ce4-a165-1b17daf37a7a.xlsx"
wb1 = openpyxl.load_workbook(path1, data_only=True)
print("=" * 80)
print("FILE 1: RATIO_등록_현황_1월_6월")
print("=" * 80)
print("Sheets:", wb1.sheetnames)

for sheet_name in wb1.sheetnames:
    ws = wb1[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ---")
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(40, ws.max_row), values_only=True), 1):
        if all(v is None for v in row):
            continue
        print(f"  R{row_idx}: {list(row)}")
    if ws.max_row > 40:
        print(f"  ... (total {ws.max_row} rows)")

print("\n\n")

# File 2: CAD 의뢰 현황
path2 = r"C:\Users\user\.openclaw\media\inbound\CAD_의뢰_현황_1월_6월---da0cda50-0f9b-4cbb-953b-36a1e17c89fc.xlsx"
wb2 = openpyxl.load_workbook(path2, data_only=True)
print("=" * 80)
print("FILE 2: CAD_의뢰_현황_1월_6월")
print("=" * 80)
print("Sheets:", wb2.sheetnames)

for sheet_name in wb2.sheetnames:
    ws = wb2[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols) ---")
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(40, ws.max_row), values_only=True), 1):
        if all(v is None for v in row):
            continue
        print(f"  R{row_idx}: {list(row)}")
    if ws.max_row > 40:
        print(f"  ... (total {ws.max_row} rows)")
