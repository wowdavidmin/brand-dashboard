import openpyxl

# Check the raw data file
raw_paths = [
    r'C:\Users\user\.openclaw\media\inbound\브랜드별_직접원가_raw_v1.0_20260629---47ea1d42-8317-4a51-a9fa-d5c2e2764f59.xlsx',
    r'C:\Users\user\.openclaw\media\inbound\브랜드별_직접원가_raw_v1.0_20260629---b99ec911-9409-4c91-ae86-0a63ab160c8c.xlsx',
]

for path in raw_paths:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        fname = path.split('---')[0].rsplit(chr(92), 1)[-1]
        print(f'=== {fname} ===')
        print(f'  Sheets: {wb.sheetnames}')
        print(f'  Rows: {ws.max_row}, Cols: {ws.max_column}')
        # Print all headers
        headers = []
        for c in range(1, ws.max_column+1):
            v = ws.cell(1, c).value
            if v is not None:
                headers.append(f'C{c}={v}')
        print(f'  Headers: {", ".join(headers)}')
        
        # Count unique 부서 values
        dept_col = None
        brand_col = None
        season_col = None
        qty_col = None
        price_col = None
        cm_col = None
        for c in range(1, ws.max_column+1):
            h = str(ws.cell(1, c).value or '')
            hl = h.lower()
            if '부서' in h:
                dept_col = c
            if '브랜드' in h or h == 'brand' or h.lower() == 'brand':
                brand_col = c
            if '시즌' in h:
                season_col = c
            if 'order' in hl or '오더' in h or '수량' in h or 'qty' in hl:
                qty_col = c
            if '단가' in h or 'price' in hl or '매출' in h:
                price_col = c
            if 'cm' in hl or '마진' in h:
                cm_col = c

        print(f'  부서(col): {dept_col}, 브랜드(col): {brand_col}, 시즌(col): {season_col}')
        print(f'  수량(col): {qty_col}, 단가/매출(col): {price_col}, CM(col): {cm_col}')

        # Print sample data
        if dept_col:
            depts = set()
            for r in range(2, min(ws.max_row+1, 100)):
                v = ws.cell(r, dept_col).value
                if v:
                    depts.add(str(v))
            print(f'  부서 values (sample): {sorted(depts)[:20]}')
        
        if brand_col:
            brands = set()
            for r in range(2, min(ws.max_row+1, 100)):
                v = ws.cell(r, brand_col).value
                if v:
                    brands.add(str(v))
            print(f'  브랜드 values (sample): {sorted(brands)[:20]}')
        
        print()
        wb.close()
    except Exception as e:
        import traceback
        print(f'Error: {e}')
        traceback.print_exc()
