"""Check v13 ORDER total vs raw data PO total"""
import openpyxl

# v13 file
v13 = r'C:\업무자료_약진&JS\업무 인수\영업데이터파일\2026년 6월\4th week\영업데이터파일_v13_20260702.xlsx'
wb = openpyxl.load_workbook(v13)
ws = wb.worksheets[0]

# ORDER Total row (row 10), sum S.TOTAL columns
stotal_cols = [8, 11, 14, 17, 20, 23, 25, 30, 32, 34]
v13_total = 0
for c in stotal_cols:
    v = ws.cell(10, c).value
    print(f'  C{c} ({ws.cell(5, c).value}): {v:>12,}')
    v13_total += float(v if v not in ('-', None, '') else 0)
print(f'  {"v13 ORDER Total":>25s}: {v13_total:>12,.0f}')

# Also check raw data total  
print(f'\nExpected (raw data PO): 121,176,138')
print(f'Difference: {v13_total - 121176138:>+,.0f}')
wb.close()
