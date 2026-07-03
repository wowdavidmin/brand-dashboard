import openpyxl
from collections import Counter

# Check the raw data for all 부서 values
raw_path = r'C:\Users\user\.openclaw\media\inbound\브랜드별_직접원가_raw_v1.0_20260629---47ea1d42-8317-4a51-a9fa-d5c2e2764f59.xlsx'
wb = openpyxl.load_workbook(raw_path, read_only=True, data_only=True)
ws = wb.worksheets[0]

# Collect all unique 부서 values
depts = Counter()
brands_by_dept = {}

for r in range(2, ws.max_row+1):
    dept = ws.cell(r, 5).value
    brand = ws.cell(r, 4).value
    season = ws.cell(r, 11).value
    
    if dept:
        dept_str = str(dept).strip()
        depts[dept_str] += 1
        if dept_str not in brands_by_dept:
            brands_by_dept[dept_str] = set()
        if brand:
            brands_by_dept[dept_str].add(str(brand).strip())

print('=== Departments in raw data ===')
for d, cnt in sorted(depts.items()):
    brands = sorted(brands_by_dept.get(d, []))
    print(f'{d}: {cnt} rows, brands: {brands}')

wb.close()
print()

# Now check which 부서 = BU3 in the v12 mapping
print('=== v12 file BU mapping ===')
v12 = r'C:\Users\user\.openclaw\media\inbound\영업데이터파일_v12_20260702---c5b345bd-47c5-47cf-b54c-e337523c2574.xlsx'
wb2 = openpyxl.load_workbook(v12)
ws2 = wb2.worksheets[0]

# Row 4 (header row for ORDER)
print('BU headers (Row 4):')
for c in range(2, 35):
    v = ws2.cell(4, c).value
    if v:
        print(f'  C{c}: {v}')

print()
print('Brand headers (Row 5):')
for c in range(2, 35):
    v = ws2.cell(5, c).value
    if v:
        print(f'  C{c}: {v}')
