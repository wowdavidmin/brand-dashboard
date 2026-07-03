import openpyxl

# Current v12 file
src = r'C:\Users\user\.openclaw\workspace\영업데이터파일_v12_20260702_processed.xlsx'
wb = openpyxl.load_workbook(src)
ws = wb.worksheets[0]

print('=== ORDER 수량 - BU3 (cols 15-20) ===')
for r in [6,7,8,9,10]:
    season = ws.cell(r, 2).value
    c15 = ws.cell(r, 15).value
    c16 = ws.cell(r, 16).value
    c17 = ws.cell(r, 17).value
    c18 = ws.cell(r, 18).value
    c19 = ws.cell(r, 19).value
    c20 = ws.cell(r, 20).value
    print(f'  {season}: BR={c15}, IPSS ON={c16}, S.TOTAL_1부={c17} | KOHLS={c18}, URBAN={c19}, S.TOTAL_2부={c20}')

print()
print('=== 총매출액 - BU3 ===')
for r in [15,16,17,18,19]:
    season = ws.cell(r, 2).value
    print(f'  {season}: BR={ws.cell(r,15).value}, IPSS ON={ws.cell(r,16).value}, S.TOTAL_1부={ws.cell(r,17).value} | KOHLS={ws.cell(r,18).value}, URBAN={ws.cell(r,19).value}, S.TOTAL_2부={ws.cell(r,20).value}')

print()
print('=== 순매출액 - BU3 ===')
for r in [24,25,26,27,28]:
    season = ws.cell(r, 2).value
    print(f'  {season}: BR={ws.cell(r,15).value}, IPSS ON={ws.cell(r,16).value}, S.TOTAL_1부={ws.cell(r,17).value} | KOHLS={ws.cell(r,18).value}, URBAN={ws.cell(r,19).value}, S.TOTAL_2부={ws.cell(r,20).value}')
