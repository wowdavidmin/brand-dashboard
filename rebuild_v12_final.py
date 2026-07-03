import pandas as pd
import numpy as np
import openpyxl

# Read raw data
raw_path = r'C:\Users\user\.openclaw\media\inbound\브랜드별_직접원가_raw_v1.0_20260629---47ea1d42-8317-4a51-a9fa-d5c2e2764f59.xlsx'

df = pd.read_excel(raw_path, usecols=[2, 4, 10, 23, 24, 25, 40], header=0)
df.columns = ['바이어', '부서코드', '시즌', 'PO수량', 'PO금액', '매출할인', 'CMT계']

# Filter to 26 season
df_26 = df[df['시즌'].str.startswith('26 ', na=False)].copy()
print(f'Total rows (26 season): {len(df_26)}')

# Aggregation
agg = df_26.groupby(['부서코드', '바이어', '시즌']).agg({
    'PO수량': 'sum', 'PO금액': 'sum', '매출할인': 'sum', 'CMT계': 'sum'
}).reset_index()

# === BU3 DETAIL ===
for dept in ['BU3_D1', 'BU3_D2', 'BU3_D3', 'BU3_D4']:
    sub = agg[agg['부서코드'] == dept]
    print(f'\n=== {dept} ===')
    for _, row in sub.iterrows():
        net = row['PO금액'] - row['매출할인']
        avg_p = row['PO금액'] / row['PO수량'] if row['PO수량'] > 0 else 0
        cm_u = row['CMT계'] / row['PO수량'] if row['PO수량'] > 0 else 0
        print(f'  {row["바이어"]:25s} {row["시즌"]:8s} 수량={row["PO수량"]:>10.0f} 매출={row["PO금액"]:>12.2f} 할인={row["매출할인"]:>10.2f} 단가={avg_p:.2f} CM={cm_u:.2f}')

# === COMPARE v12 vs Raw ===
print('\n\n=== v12 vs Raw: BU3 ORDER 수량 ===')
v12_path = r'C:\Users\user\.openclaw\media\inbound\영업데이터파일_v12_20260702---c5b345bd-47c5-47cf-b54c-e337523c2574.xlsx'
wb = openpyxl.load_workbook(v12_path)
ws = wb.worksheets[0]

# Column in v12 -> brand name in raw
v12_cols = [
    (15, 'BANANA REPUBLIC', 'BU3_D1'),
    (16, 'ON IPSS', 'BU3_D1'),
    (18, "KOHL`S", 'BU3_D2'),
    (19, 'URBAN OUTFITTERS', 'BU3_D2'),
    (21, 'AERIE', 'BU3_D3'),
    (22, 'Offline', 'BU3_D3'),
    (24, 'MLB', 'BU3_D4'),
]

season_rows = {6: '26 SPR', 7: '26 SUM', 8: '26 FAL', 9: '26 HOL', 10: 'Total'}
total_rows = {15: '26 SPR', 16: '26 SUM', 17: '26 FAL', 18: '26 HOL', 19: 'Total'}
net_rows = {24: '26 SPR', 25: '26 SUM', 26: '26 FAL', 27: '26 HOL', 28: 'Total'}

for col, brand, dept in v12_cols:
    print(f'\n--- {brand} (C{col}, {dept}) ---')
    for row, season in season_rows.items():
        v12_val = ws.cell(row, col).value
        v12_val = float(v12_val if v12_val not in ('-', None, '') else 0)
        
        raw_row = agg[(agg['바이어']==brand) & (agg['시즌']==season)]
        if len(raw_row) > 0:
            raw_qty = raw_row['PO수량'].values[0]
            raw_amt = raw_row['PO금액'].values[0]
            qty_diff = v12_val - raw_qty
            print(f'  {season:8s} v12_qty={v12_val:>10.0f} raw_qty={raw_qty:>10.0f} diff={qty_diff:>+10.0f} | v12_amt={ws.cell(total_rows.get(row,0),col).value if total_rows.get(row) else "?":>8} raw_amt={raw_amt/1000:>8.1f}K')
        else:
            print(f'  {season:8s} v12_qty={v12_val:>10.0f} raw=NOT_FOUND')
